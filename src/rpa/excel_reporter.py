"""
RPA Excel Reporter Module
Generates beautifully formatted multi-sheet Excel reports with conditional formatting,
clean headers, and automated portfolio / FX / backtest summaries.
Builds real-world data processing & RPA automation skills.
"""

from datetime import datetime
import os
from typing import Dict, Any, List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


class QuantExcelReporter:
    """Automates professional Excel report generation for daily quant investment tracking."""

    def __init__(self, output_dir: str = "data/reports"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_daily_report(
        self,
        portfolio_df: pd.DataFrame,
        fx_stats: Dict[str, Any],
        backtest_df: pd.DataFrame,
        filename: str = None,
        ai_briefing: Dict[str, Any] = None
    ) -> str:
        """Create a multi-tab formatted Excel workbook."""
        today_str = datetime.now().strftime("%Y%m%d")
        if not filename:
            filename = f"ISA_퀀트투자_일일리포트_{today_str}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style Definitions
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Malgun Gothic", size=14, bold=True, color="1F4E78")
        bold_font = Font(name="Malgun Gothic", size=10, bold=True)
        regular_font = Font(name="Malgun Gothic", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        # ----------------- Sheet 1: 포트폴리오 일일 현황 -----------------
        ws1 = wb.create_sheet(title="포트폴리오_현황")
        ws1.views.sheetView[0].showGridLines = True

        # Title
        ws1.merge_cells("A1:G1")
        ws1["A1"] = f"ISA 포트폴리오 일일 정산 리포트 ({datetime.now().strftime('%Y-%m-%d')})"
        ws1["A1"].font = title_font
        ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")

        # Table Headers
        headers1 = list(portfolio_df.columns)
        ws1.append([])  # Row 2 empty
        ws1.append(headers1)  # Row 3

        for col_idx, _ in enumerate(headers1, 1):
            cell = ws1.cell(row=3, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Add Data rows
        for r_idx, row in portfolio_df.iterrows():
            ws1.append(row.tolist())
            curr_row = ws1.max_row
            for col_idx in range(1, len(headers1) + 1):
                cell = ws1.cell(row=curr_row, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                cell_str = str(cell.value or "")

                if "%" in cell_str:
                    cell.alignment = right_align
                    try:
                        val = float(cell_str.replace("%", "").replace(",", "").replace("+", "").strip())
                        if val > 0:
                            cell.font = Font(name="Malgun Gothic", size=10, bold=True, color="C00000")
                        elif val < 0:
                            cell.font = Font(name="Malgun Gothic", size=10, bold=True, color="0070C0")
                    except ValueError:
                        pass
                elif any(char.isdigit() for char in cell_str) and ("원" in cell_str or "주" in cell_str or isinstance(cell.value, (int, float))):
                    cell.alignment = right_align
                else:
                    cell.alignment = center_align

        # ----------------- Sheet 2: 환율 & 기초자산 분석 -----------------
        ws2 = wb.create_sheet(title="환율_상관관계_분석")
        ws2.views.sheetView[0].showGridLines = True

        ws2.merge_cells("A1:D1")
        ws2["A1"] = "기초 자산 vs 원/달러 환율 상관관계 & 방어력 분석 요약"
        ws2["A1"].font = title_font

        ws2.append([])
        stats_headers = ["지표 항목", "값", "단위 / 설명"]
        ws2.append(stats_headers)
        for col_idx in range(1, 4):
            c = ws2.cell(row=3, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align

        stat_rows = [
            ["분석 기간", f"{fx_stats.get('period_start')} ~ {fx_stats.get('period_end')}", "영업일 기준"],
            ["미국 지수 vs 환율 상관계수", fx_stats.get("correlation_us_fx"), "-1에 가까울수록 환율 방어력 우수"],
            ["국내 ETF 누적 수익률", f"{fx_stats.get('total_kr_etf_return_pct'):+.2f}%", "실제 투자 수익률"],
            ["미국 원본 ETF 수익률", f"{fx_stats.get('total_us_etf_return_pct'):+.2f}%", "달러 기준"],
            ["원/달러 환율 변동률", f"{fx_stats.get('total_fx_return_pct'):+.2f}%", "환차익/환차손 기여도"],
            ["이론적 환노출 수익률", f"{fx_stats.get('theoretical_krw_return_pct'):+.2f}%", "(1+주가)*(1+환율)-1"],
            ["평균 괴리율", f"{fx_stats.get('disparity_mean_pct'):+.2f}%", "이론가 대비 실제 ETF 가격 편차"],
            ["최근 괴리율", f"{fx_stats.get('disparity_latest_pct'):+.2f}%", "최근일 기준"],
            ["미국 증시 하락일 환율 상승 방어 확률", f"{fx_stats.get('fx_defense_probability_pct')}%", "미국장 하락 시 달러 상승 방어 빈도"],
        ]

        for item in stat_rows:
            ws2.append(item)
            curr_row = ws2.max_row
            for col_idx in range(1, 4):
                c = ws2.cell(row=curr_row, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                c.alignment = center_align if col_idx != 3 else Alignment(horizontal="left", vertical="center")

        # ----------------- Sheet 3: 백테스트 비교 -----------------
        ws3 = wb.create_sheet(title="자산배분_백테스트")
        ws3.views.sheetView[0].showGridLines = True

        ws3.merge_cells("A1:G1")
        ws3["A1"] = "포트폴리오 비중 조절 (7:3 vs 5:5) 과거 백테스트 결과"
        ws3["A1"].font = title_font
        ws3.append([])

        bt_rows = dataframe_to_rows(backtest_df, index=False, header=True)
        for r_idx, row in enumerate(bt_rows, 3):
            ws3.append(row)
            for col_idx in range(1, len(row) + 1):
                c = ws3.cell(row=r_idx, column=col_idx)
                c.border = thin_border
                if r_idx == 3:
                    c.fill = header_fill
                    c.font = header_font
                    c.alignment = center_align
                else:
                    c.font = regular_font
                    c.alignment = center_align

        # ----------------- Sheet 4: AI 스마트 시황 브리핑 -----------------
        if ai_briefing:
            ws4 = wb.create_sheet(title="AI_시황브리핑")
            ws4.views.sheetView[0].showGridLines = True

            ws4.merge_cells("A1:G1")
            ws4["A1"] = f"{ai_briefing.get('title', 'AI 스마트 퀀트 시황 브리핑')} ({ai_briefing.get('generated_time', '')})"
            ws4["A1"].font = title_font

            ws4.merge_cells("A2:G2")
            ws4["A2"] = f"총괄 진단: {ai_briefing.get('summary_message', '')}"
            ws4["A2"].font = bold_font
            ws4["A2"].alignment = Alignment(horizontal="left", vertical="center")

            ai_headers = ["종목코드", "종목명", "계좌구분", "현재가", "AI 진단상태", "예상 밴드/시초가", "AI 맞춤 행동 가이드"]
            ws4.append([])  # Row 3 empty
            ws4.append(ai_headers)  # Row 4

            for col_idx in range(1, len(ai_headers) + 1):
                c = ws4.cell(row=4, column=col_idx)
                c.fill = header_fill
                c.font = header_font
                c.alignment = center_align

            for s in ai_briefing.get("stock_reports", []):
                price_str = f"{s.get('current_price', 0):,.0f}원"
                if ai_briefing.get("mode") == "morning":
                    status_str = s.get("status", "")
                    range_str = f"{s.get('range_low', 0):,.0f}원 ~ {s.get('range_high', 0):,.0f}원 (하위 {s.get('percentile_pos', 50):.0f}%)"
                else:
                    status_str = s.get("direction_badge", "")
                    range_str = f"예상시초가 {s.get('est_open_price', 0):,.0f}원 ({s.get('predicted_gap_pct', 0):+.2f}%)"

                row_vals = [
                    s.get("code", ""),
                    s.get("name", ""),
                    s.get("account", ""),
                    price_str,
                    status_str,
                    range_str,
                    s.get("action_comment", "")
                ]
                ws4.append(row_vals)
                curr_row = ws4.max_row
                for col_idx in range(1, len(row_vals) + 1):
                    c = ws4.cell(row=curr_row, column=col_idx)
                    c.font = regular_font
                    c.border = thin_border
                    if col_idx in [1, 3, 5]:
                        c.alignment = center_align
                    elif col_idx in [4, 6]:
                        c.alignment = right_align
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for ws in wb.worksheets:
            for col in ws.columns:
                col_idx = col[0].column
                col_letter = get_column_letter(col_idx)
                # Ignore row 1 and 2 merged titles
                data_cells = [cell for cell in col if cell.row > 2]
                max_len = max((len(str(cell.value or '')) for cell in data_cells), default=10)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 60)

        wb.save(filepath)
        print(f"[RPA] Report generated successfully: {filepath}")
        return filepath
